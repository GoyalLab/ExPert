#!/usr/bin/env python
# SPDX-License-Identifier: MIT
#
# ExPert: a contrastive perturbation-prediction model for single-cell data.
#
# Copyright (c) 2024-2026 Leon Schwartz
# Northwestern University, Goyal Lab
#
# This file is part of the ExPert package and is released under the MIT
# License. See the LICENSE file in the project root for the full text.

"""Fig. 6I: flow cytometry of the sorted-cell cohort -- monocyte mitochondrial mass (MitoTracker) and cellular ROS
(CellROX), each as the % change of the geometric MFI, stimulated vs unstimulated, per patient; PMD vs control.

Source: the 'Stim vs Unstim % change' sheet of the collaborators' workbook (--xlsx). Rows the collaborators
highlighted in that sheet are excluded (read from the cell fill, not hard-coded) and listed in the source table.
Test: two-sided Student t-test (reproduces the collaborators' prelim p-values); Welch and exact Mann-Whitney p in the
stats table.

    PYTHONPATH=paper/scripts python paper/scripts/fig6/fig6_flow_plot.py --xlsx ... --out-dir ... --src-dir ... --vizstyle-root ...
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt  # noqa: E402
plt.rcParams.update({'svg.fonttype': 'none', 'pdf.fonttype': 42})

PMD = {'43', '54', '63', '67', '75'}
# (block title in the sheet, column header, y label)
READOUTS = [('Mono MT-TMRE', 'gMFI (MT) % change', 'MitoTracker'),
            ('Mono ROS', 'gMFI (CellRox) % change', 'CellROX')]


def read_block(xlsx, title, column):
    """One block of the % change sheet: patient, value (%), highlighted (= excluded by the collaborators)."""
    import openpyxl
    ws = openpyxl.load_workbook(xlsx, data_only=True)['Stim vs Unstim % change']
    wf = openpyxl.load_workbook(xlsx)['Stim vs Unstim % change']     # fills survive either way; keep it explicit
    rows = list(ws.iter_rows())
    start = next(i for i, r in enumerate(rows) if r[0].value == title)
    head = [c.value for c in rows[start + 1]]
    col = head.index(column)
    out = []
    for r in rows[start + 2:]:
        name = r[0].value
        if not isinstance(name, str) or not name.startswith('Patient'):
            break
        v = r[col].value
        fill = wf[r[0].coordinate].fill
        hl = bool(fill and fill.fill_type and fill.fgColor.rgb not in (None, '00000000'))
        pid = name.split()[1]
        out.append(dict(readout=title, patient=f'P{pid}', group='PMD' if pid in PMD else 'control',
                        pct=100 * float(v) if isinstance(v, (int, float)) else np.nan, excluded=hl))
    return pd.DataFrame(out)


def main(a):
    global VS
    sys.path.insert(0, str(Path(a.vizstyle_root)))
    from src import vizstyle as VS
    from scipy.stats import mannwhitneyu, ttest_ind
    VS.apply()
    plt.rcParams.update({'figure.facecolor': 'white', 'axes.facecolor': 'white', 'savefig.facecolor': 'white'})
    out = Path(a.out_dir)
    T = pd.concat([read_block(a.xlsx, t, c) for t, c, _ in READOUTS], ignore_index=True)
    col = {'PMD': VS.C['blue'], 'control': '#a9a8a2'}
    fig, axs = plt.subplots(2, 1, figsize=(a.width, a.height), gridspec_kw={'hspace': 0.55})
    rng = np.random.default_rng(0)
    stats = []
    for ax, (title, _, lab) in zip(axs, READOUTS):
        x = T[(T.readout == title) & ~T.excluded].dropna(subset=['pct'])
        for k, g in enumerate(('control', 'PMD')):
            v = x[x.group == g].pct.to_numpy()
            ax.scatter(k + rng.uniform(-0.08, 0.08, len(v)), v, s=36, color=col[g], edgecolor='white', lw=0.6,
                       zorder=3)
            ax.plot([k - 0.22, k + 0.22], [v.mean()] * 2, color=col[g], lw=2.2, zorder=4)   # mean; black = DEG elsewhere
        c_, p_ = x[x.group == 'control'].pct, x[x.group == 'PMD'].pct
        pt = ttest_ind(p_, c_).pvalue
        pw = ttest_ind(p_, c_, equal_var=False).pvalue
        pm = mannwhitneyu(p_, c_, alternative='two-sided', method='exact').pvalue
        stats.append(dict(readout=lab, n_control=len(c_), n_pmd=len(p_), mean_control=c_.mean(), mean_pmd=p_.mean(),
                          p_student=pt, p_welch=pw, p_mannwhitney=pm))
        top = max(x.pct.max(), 0)
        span = x.pct.max() - min(x.pct.min(), 0)
        yb = top + 0.12 * span
        ax.plot([0, 0, 1, 1], [yb - 0.04 * span, yb, yb, yb - 0.04 * span], color=VS.INK2, lw=1.0)
        ax.text(0.5, yb + 0.03 * span, f'p = {pt:.2f}', ha='center', va='bottom', fontsize=VS.FONT_MIN - 2,
                color=VS.INK2)
        ax.axhline(0, color=VS.MUTED, lw=0.8, zorder=0)
        ax.set_xticks([0, 1]); ax.set_xticklabels(['control', 'PMD'])
        ax.set_xlim(-0.5, 1.5)
        ax.set_ylim(min(x.pct.min(), 0) - 0.08 * span, yb + 0.2 * span)
        ax.set_ylabel(f'{lab}\nΔ MFI, stim (%)')
        ax.yaxis.set_major_locator(matplotlib.ticker.MaxNLocator(nbins=4))
        for s_ in ('top', 'right'):
            ax.spines[s_].set_visible(False)
        ax.grid(False)
    VS.save(fig, out, a.name)
    Path(a.src_dir).mkdir(parents=True, exist_ok=True)
    T.to_csv(Path(a.src_dir) / f'{a.name}_source.csv', index=False)
    S = pd.DataFrame(stats)
    S.to_csv(Path(a.src_dir) / f'{a.name}_stats.csv', index=False)
    print(S.round(4).to_string(index=False))
    print('excluded:', T[T.excluded].groupby('readout').patient.apply(list).to_dict())


def build_parser():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--xlsx', required=True, help='mito_flow_combined_raw.xlsx')
    ap.add_argument('--out-dir', required=True)
    ap.add_argument('--src-dir', required=True,
                    help='source tables behind the panels (paper/figures/<fig>/data)')
    ap.add_argument('--vizstyle-root', required=True)
    ap.add_argument('--name', default='fig6i_flow')
    ap.add_argument('--width', type=float, default=2.6, help='slim panel width (in)')
    ap.add_argument('--height', type=float, default=6.0)
    return ap


if __name__ == '__main__':
    main(build_parser().parse_args())
